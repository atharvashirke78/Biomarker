#!/usr/bin/env python3
"""Build Hosseini2023_vs_Pipeline_comparison.xlsx (formulas) and Hosseini2023_vs_Pipeline_comparison_values.xlsx (same content, static values).
Read-only on every existing file. Run:
  PYTHONPATH=/workspace/scratch_verify/pylib /workspace/Biomarker_pipeline/env/prefix/pyqc/bin/python build_comparison.py
Every number is computed from the files listed in INPUTS (printed at the start and written to sheet Method_notes)."""
import csv, json, math, os, re, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

B = "/workspace/Biomarker_pipeline"; R = f"{B}/results"; V = f"{R}/09_verification"; X = f"{V}/input/MOESM2_sheets"
OUT = f"{V}/step11_comparison_workbook"; XLSX = f"{OUT}/Hosseini2023_vs_Pipeline_comparison.xlsx"; XLSX_VAL = f"{OUT}/Hosseini2023_vs_Pipeline_comparison_values.xlsx"
PADJ, LFC = 0.05, 1.0
INPUTS = {
    "paper S1 (samples)": f"{X}/Table_S1.tsv", "paper S2 (read counts)": f"{X}/Table_S2.tsv",
    "paper S3 (DEGs, up/down)": f"{V}/input/hosseini_2023_TableS3_genes.tsv",
    "paper S4 GO BP down": f"{X}/Table_S4_Biological_process_.tsv", "paper S5 GO BP up": f"{X}/Table_S5_Biological_process_.tsv",
    "paper S6 KEGG down": f"{X}/Table_S6_Downregulated_.tsv", "paper S6 KEGG up": f"{X}/Table_S6_Upregulated_.tsv",
    "ours DESeq2 featureCounts (primary)": f"{R}/06_dge/deseq2_results_all.tsv", "ours DESeq2 Salmon (secondary)": f"{R}/06_dge/deseq2_salmon_results_all.tsv",
    "ours DESeq2 summary": f"{R}/06_dge/deseq2_summary.json", "ours normalized counts (sample list)": f"{R}/06_dge/deseq2_normalized_counts.tsv",
    "ours featureCounts count matrix (19 samples)": f"{R}/04_counts/counts_featurecounts.tsv", "ours gene annotation (id, symbol, biotype)": f"{B}/ref/gene_id_to_symbol.tsv",
    "ours ORA GO BP up": f"{R}/07_pathway/ora/ora_up_GO_BP.tsv", "ours ORA GO BP down": f"{R}/07_pathway/ora/ora_down_GO_BP.tsv",
    "ours ORA KEGG up": f"{R}/07_pathway/ora/ora_up_KEGG.tsv", "ours ORA KEGG down": f"{R}/07_pathway/ora/ora_down_KEGG.tsv",
    "ours sample map": f"{R}/sample_map.tsv", "ours fastp JSON (per run)": f"{R}/01_fastp/<SRR>/fastp.json",
    "step4b leave-one-out T2_B": f"{V}/step9_paper_comparison/step4b_deseq2_without_T2_B_results.tsv",
    "step8 refit without N4_A,N5_A": f"{V}/step8_extra_questions/deseq2_without_N4A_N5A_results.tsv",
    "tool versions": f"{B}/docs/tool_versions.tsv",
}
print("INPUT FILES USED:"); [print(f"  {k:46s} {v}") for k, v in INPUTS.items()]

# ---------------------------------------------------------------- helpers
def norm(x): return re.sub(r"[^a-z0-9]+", " ", str(x).lower()).strip()
def num(x): return float(re.sub(r"[^0-9.]", "", str(x).split("(")[0])) if str(x).strip() else float("nan")
def read_enrichr(f):
    d = pd.read_csv(f, sep="\t", skiprows=1, index_col=0, quoting=csv.QUOTE_NONE, dtype=str)
    d["padj"] = d["Adjusted.P.value"].astype(float); return d
def lchoose(n, k): return math.lgamma(n + 1) - math.lgamma(k + 1) - math.lgamma(n - k + 1)
def hyper_sf(k, K, n, N):          # P(X >= k), X ~ Hypergeom(N, K, n), summed in log space
    ls = [lchoose(K, i) + lchoose(N - K, n - i) - lchoose(N, n) for i in range(k, min(K, n) + 1)]
    m = max(ls); return math.exp(m) * sum(math.exp(l - m) for l in ls) if ls else 0.0
def ratio(a, b): return "n/a" if (not isinstance(b, (int, float)) or b == 0) else a / b

def load_de(path):
    d = pd.read_csv(path, sep="\t"); d = d[d.gene_name.notna() & (d.gene_name != "")].copy()
    d["_ok"] = d.padj.notna(); d = d.sort_values(["_ok", "baseMean"], ascending=[False, False]).drop_duplicates("gene_name")   # one row per symbol (as step5)
    d["sig"] = d.padj.notna() & (d.padj < PADJ) & (d.log2FoldChange.abs() >= LFC)
    d["dir"] = d.log2FoldChange.map(lambda v: "" if pd.isna(v) else ("up" if v > 0 else "down"))
    return d.set_index("gene_name")

fc = load_de(INPUTS["ours DESeq2 featureCounts (primary)"]); sal = load_de(INPUTS["ours DESeq2 Salmon (secondary)"])
s3 = pd.read_csv(INPUTS["paper S3 (DEGs, up/down)"], sep="\t"); s3["gene"] = s3.gene.str.strip()
assert not s3.gene.duplicated().any(), "duplicate symbol in S3"
pdir = dict(zip(s3.gene, s3.direction_in_cancer)); paper = {d: set(s3.gene[s3.direction_in_cancer == d]) for d in ("up", "down")}; s3i = s3.set_index("gene")

CAT_M, CAT_O = "Matched (same direction)", "Opposite direction"
CAT_P = ["Paper only: tested, padj >= 0.05", "Paper only: padj < 0.05 but |log2FC| < 1", "Paper only: padj NA (Cook's / low count)", "Paper only: not in our DESeq2 table (filtered or not annotated)"]
CAT_X = "Ours only (not in paper list)"
def category(g, d, tab):
    if g not in tab.index: return CAT_P[3]
    r = tab.loc[g]
    if r.sig: return CAT_M if r["dir"] == d else CAT_O
    if pd.isna(r.padj): return CAT_P[2]
    if r.padj >= PADJ: return CAT_P[0]
    return CAT_P[1]

def gene_sets(tab):
    out = {}; cats = {g: category(g, d, tab) for g, d in pdir.items()}
    for d in ("up", "down"):
        pl = paper[d]; ours = set(tab.index[tab.sig & (tab["dir"] == d)])
        m = sum(cats[g] == CAT_M for g in pl); o = sum(cats[g] == CAT_O for g in pl)
        opp_ours = len(ours & (set(pdir) - pl)); extra = len(ours - set(pdir))
        out[d] = dict(paper=len(pl), ours=len(ours), matched=m, opposite=o, missed=len(pl) - m - o, opposite_ours_side=opp_ours, extra=extra)
        assert len(ours) == m + opp_ours + extra, (d, len(ours), m, opp_ours, extra)
    return out, cats
GS_fc, CATS_fc = gene_sets(fc); GS_sal, CATS_sal = gene_sets(sal)

tested = fc[fc.padj.notna()]; N = len(tested); ENR = {}
for d in ("up", "down"):
    ours = set(tested.index[tested.sig & (tested["dir"] == d)]); pl = paper[d] & set(tested.index); k = len(ours & pl)
    exp = len(ours) * len(pl) / N; ENR[d] = dict(ours=len(ours), paper_in_universe=len(pl), overlap=k, expected=exp, fold=k / exp, p=hyper_sf(k, len(pl), len(ours), N))

# ---------------------------------------------------------------- read counts / samples
ss = pd.read_csv(INPUTS["ours sample map"], sep="\t")
s2 = {}
for line in open(INPUTS["paper S2 (read counts)"]):
    c = line.rstrip("\n").split("\t")
    if re.match(r"^SRR\d+$", c[0]): s2[c[0]] = num(c[1])
rc = []
for _, r in ss.iterrows():
    j = json.load(open(f"{R}/01_fastp/{r.srr}/fastp.json")); ours_reads = j["summary"]["before_filtering"]["total_reads"] / 2
    rc.append([r.srr, s2.get(r.srr), ours_reads, bool(s2.get(r.srr) == ours_reads), r.label, r.study, r.group, "Excluded from DE analysis (R2 is a copy of R1)" if r.label == "T5_B" else ""])
n_rc_match = sum(x[3] for x in rc)
cols_used = list(pd.read_csv(INPUTS["ours normalized counts (sample list)"], sep="\t", nrows=1).columns[1:])
used = ss[ss.srr.isin(cols_used)]; n_t = int((used.group == "tumour").sum()); n_n = int((used.group == "normal").sum())
paper_total = int(re.search(r"Total Samples:\s*(\d+)", open(INPUTS["paper S1 (samples)"]).read()).group(1))

# ---------------------------------------------------------------- pathways (opposite direction included)
def pathway_rows():
    rows = []; counts = {}; P = {}; O = {}
    for typ in ("GO", "KEGG"):
        for d in ("down", "up"):
            f = INPUTS[f"ours ORA {'GO BP' if typ == 'GO' else 'KEGG'} {d}"]; o = pd.read_csv(f, sep="\t"); o = o[o["p.adjust"] < PADJ].copy()
            if typ == "GO":
                p = read_enrichr(INPUTS["paper S4 GO BP down" if d == "down" else "paper S5 GO BP up"]); p = p[p.padj < PADJ].copy()
                p["id"] = p.Term.str.extract(r"\((GO:\d+)\)")[0]; p["name"] = p.Term.str.replace(r"\s*\(GO:\d+\)\s*$", "", regex=True); p["key"] = p["id"]; o["key"] = o["ID"]
            else:
                p = read_enrichr(INPUTS["paper S6 KEGG down" if d == "down" else "paper S6 KEGG up"]); p = p[p.padj < PADJ].copy()
                p["id"] = ""; p["name"] = p.Term; p["key"] = p.name.map(norm); o["key"] = o.Description.map(norm)
            P[(typ, d)] = p; O[(typ, d)] = o
    for typ in ("GO", "KEGG"):
        for d in ("down", "up"):
            e = "up" if d == "down" else "down"; p, o = P[(typ, d)], O[(typ, d)]; pk, ok = set(p.key), set(o.key); pk_e, ok_e = set(P[(typ, e)].key), set(O[(typ, e)].key)
            omap = o.set_index("key"); omap_e = O[(typ, e)].set_index("key")
            matched = pk & ok; opp = (pk & ok_e) - ok; extra = ok - pk - pk_e
            counts[(typ, d)] = dict(paper=len(pk), ours=len(ok), matched=len(matched), opposite=len(opp), missed=len(pk) - len(matched) - len(opp), extra=len(extra), opposite_ours_side=len((ok & pk_e) - pk))
            for _, r in p.sort_values("padj").iterrows():
                k = r.key; note = ""
                if k in ok: st = "Yes"; ours_t, ours_p = omap.loc[k, "Description"], float(omap.loc[k, "p.adjust"])
                elif k in ok_e: st = "Opposite direction"; ours_t, ours_p = omap_e.loc[k, "Description"], float(omap_e.loc[k, "p.adjust"]); note = f"Significant in our {e} list, not in our {d} list."
                else:
                    st = "No (paper only)"; ours_t, ours_p = "", None
                    if typ == "KEGG":
                        note = "Paper only (not in our ORA table). KEGG release/name difference is a likely explanation, not verified."
                        if norm(r["name"]) in ("spliceosome", "rna transport", "mrna surveillance pathway"): note = "Paper only (confirmed: absent from our ORA table). " + note
                    else: note = "Paper only: not significant or not tested in our ORA (GO matched by ID)."
                rows.append([typ, d, r["name"], r["id"] or "", r.padj, st, ours_t, ours_p, note])
            for _, r in o.sort_values("p.adjust").iterrows():
                if r.key in extra:
                    rows.append([typ, d, "", r["ID"], None, "Ours only", r["Description"], float(r["p.adjust"]), "Ours only." + (" KEGG release difference possible, not verified." if typ == "KEGG" else "")])
    return rows, counts
PATH_ROWS, PC = pathway_rows()

# ---------------------------------------------------------------- 18 genes
G18 = ["GUCA2A", "CLCA4", "CLCA1", "AQP8", "SLC26A3", "APOA4", "PYY", "GUCA2B", "ZG16", "CD177", "COL3A1", "COL1A1", "SPARC", "TGFBI", "CXCL8", "KRT6A", "KRT6B", "KRT17"]
assert all(g in pdir for g in G18), [g for g in G18 if g not in pdir]
def other_padj(path, g):
    t = pd.read_csv(path, sep="\t"); r = t[t.gene_name == g]; return None if r.empty else float(r.iloc[0].padj)
g18 = []
for g in G18:
    r = fc.loc[g] if g in fc.index else None; sig = bool(r is not None and r.sig); caveat = ""
    if r is None: caveat = "Not testable: not in our DESeq2 table (low counts, filtered by pre-filter)"
    elif pd.isna(r.padj): caveat = "Direction only: padj NA (Cook's outlier flag; estimate driven by one or two samples)"
    elif g == "COL3A1":
        a = other_padj(INPUTS["step4b leave-one-out T2_B"], g); b = other_padj(INPUTS["step8 refit without N4_A,N5_A"], g)
        caveat = f"Borderline: padj {r.padj:.3f}; without T2_B padj {a:.3f}, without N4_A+N5_A padj {b:.3f}"
    elif not sig: caveat = f"Not significant: padj {r.padj:.2f}, |log2FC| {abs(r.log2FoldChange):.2f}"
    g18.append([g, pdir[g], (r["dir"] if r is not None else ""), (None if r is None or pd.isna(r.padj) else float(r.padj)), (None if r is None else float(r.log2FoldChange)),
                "Yes" if sig else "No", ("Yes" if sig and r["dir"] == pdir[g] else ("No" if sig else "")), caveat, (sal.loc[g, "dir"] if g in sal.index else ""), (None if g not in sal.index or pd.isna(sal.loc[g, "padj"]) else float(sal.loc[g, "padj"]))])
g18_sig = sum(x[5] == "Yes" for x in g18); g18_match = sum(x[6] == "Yes" for x in g18); g18_opp = sum(x[6] == "No" and x[5] == "Yes" for x in g18); g18_dir = sum(x[2] == x[1] for x in g18)

# ---------------------------------------------------------------- opposite-direction genes, causes for "not in our DESeq2 table"
opp_genes = [g for g in pdir if CATS_fc[g] == CAT_O]
ann = pd.read_csv(INPUTS["ours gene annotation (id, symbol, biotype)"], sep="\t"); cnt = pd.read_csv(INPUTS["ours featureCounts count matrix (19 samples)"], sep="\t", index_col=0)[cols_used]
n_ge10 = (cnt >= 10).sum(axis=1); passes = n_ge10 >= 5; tot = cnt.sum(axis=1)
assert int(passes.sum()) == json.load(open(INPUTS["ours DESeq2 summary"]))["n_genes_after_prefilter"], "pre-filter count differs from DESeq2 summary"
CAUSES = ["Removed by our low-count filter: protein-coding gene", "Removed by our low-count filter: lncRNA", "Removed by our low-count filter: other biotype (pseudogene, miRNA, etc.)",
          "No matching gene symbol in our annotation (alias / renamed symbol)", "Other (in annotation and passes the filter, but not in our table)"]
def cause(g):
    a = ann[ann.gene_name == g]
    if a.empty: return CAUSES[3], "", "", None, None
    ids = a.gene_id.tolist(); best = max(ids, key=lambda i: tot[i]); bt = a.set_index("gene_id").loc[best, "gene_type"]
    if any(passes[i] for i in ids): return CAUSES[4], ";".join(ids), bt, int(cnt.loc[best].max()), int(n_ge10[best])
    c = CAUSES[0] if bt == "protein_coding" else (CAUSES[1] if bt == "lncRNA" else CAUSES[2])
    return c, ";".join(ids), bt, int(cnt.loc[best].max()), int(n_ge10[best])
notin = [[g, pdir[g]] + list(cause(g)) for g in pdir if CATS_fc[g] == CAT_P[3]]
notin.sort(key=lambda x: (CAUSES.index(x[2]), x[1], x[0]))
cause_cnt = {c: {d: sum(1 for x in notin if x[2] == c and x[1] == d) for d in ("down", "up")} for c in CAUSES}
assert sum(sum(v.values()) for v in cause_cnt.values()) == len(notin)

# balance (paper up-heavy vs ours down-heavy)
p_up, p_dn = len(paper["up"]), len(paper["down"]); o_up, o_dn = GS_fc["up"]["ours"], GS_fc["down"]["ours"]
BAL = (f"Balance of the two lists: the paper's DEG list is up-heavy ({p_up:,} up vs {p_dn:,} down, {100 * p_up / (p_up + p_dn):.1f} % up), ours is down-heavy ({o_up:,} up vs {o_dn:,} down, {100 * o_dn / (o_up + o_dn):.1f} % down). "
       "The reason for this difference was not tested here; this is a description of the imbalance, not an explanation.")

# ---------------------------------------------------------------- checks against the known values (STOP on mismatch)
_raw = pd.read_csv(INPUTS["ours DESeq2 featureCounts (primary)"], sep="\t"); _sg = _raw.padj.notna() & (_raw.padj < PADJ) & (_raw.log2FoldChange.abs() >= LFC)
chk = [("ours significant down (gene-id level, raw file)", int((_sg & (_raw.log2FoldChange < 0)).sum()), 2060), ("ours significant up (gene-id level, raw file)", int((_sg & (_raw.log2FoldChange > 0)).sum()), 1153),
       ("ours significant down (symbol level)", GS_fc["down"]["ours"], 2060), ("ours significant up (symbol level; two Ensembl IDs share a symbol)", GS_fc["up"]["ours"], 1152),
       ("S2 read counts matched", n_rc_match, 20), ("down: significant genes in paper's down list with same direction", GS_fc["down"]["matched"], 223),
       ("down: paper list size", GS_fc["down"]["paper"], 285), ("up: our significant genes found in paper's up list (same + opposite)", GS_fc["up"]["matched"] + GS_fc["up"]["opposite"], 865),
       ("up: matched in paper's up list", GS_fc["up"]["matched"], 855), ("up: paper list size", GS_fc["up"]["paper"], 4752),
       ("enrichment down (1 decimal)", round(ENR["down"]["fold"], 1), 7.2), ("enrichment up (1 decimal)", round(ENR["up"]["fold"], 1), 3.9),
       ("KEGG down shared", PC[("KEGG", "down")]["matched"], 11), ("KEGG down paper", PC[("KEGG", "down")]["paper"], 23),
       ("KEGG up shared", PC[("KEGG", "up")]["matched"], 5), ("KEGG up paper", PC[("KEGG", "up")]["paper"], 8),
       ("opposite-direction genes", len(opp_genes), 10), ("not in our table: up", sum(1 for x in notin if x[1] == "up"), 1431), ("not in our table: down", sum(1 for x in notin if x[1] == "down"), 7)]
bad = [c for c in chk if c[1] != c[2]]
print("\nKNOWN-VALUE CHECKS:"); [print(f"  {'OK ' if c[1] == c[2] else 'DIFF'} {c[0]}: got {c[1]} expected {c[2]}") for c in chk]
print("  causes for 'not in our DESeq2 table' (down / up):"); [print(f"    {c}: {v['down']} / {v['up']}") for c, v in cause_cnt.items()]
print(f"  18 genes: significant {g18_sig}, matched {g18_match}, opposite {g18_opp}, expected direction {g18_dir}/18")
if bad: print("STOP: values differ from the known checks, workbook not written."); sys.exit(3)

# ---------------------------------------------------------------- workbook
HF = PatternFill("solid", fgColor="1F4E78"); HFONT = Font(bold=True, color="FFFFFF"); THIN = Side(style="thin", color="999999"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SCI, PCT, L2 = "0.00E+00", "0.0%", "0.00"
def header(ws, row, names, col0=1):
    for i, n in enumerate(names):
        c = ws.cell(row=row, column=col0 + i, value=n); c.font = HFONT; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX
def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x

def build(static, path):
    wb = Workbook()
    def sheet(name, first=False):
        ws = wb.active if first else wb.create_sheet(); ws.title = name; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True; return ws
    def fcell(ws, r, c, formula, value, fmt=None):          # formula in the live file, the computed value in the values copy
        cell = ws.cell(row=r, column=c, value=(value if static else formula))
        if fmt: cell.number_format = fmt
        return cell
    # 1 Summary ------------------------------------------------------------
    ws = sheet("Summary", True)
    header(ws, 1, ["Item", "Paper", "Our pipeline", "Matched", "Opposite direction", "Missed (paper only)", "Extra (ours only)", "Recall %", "Our hits also in paper list %", "Check: matched + opposite + missed = paper", "Note"])
    MS = " Ours-only genes are not necessarily false positives: the paper used metaSeq on 120 samples."
    rows = [["Samples used", f"{paper_total} (60 cancer + 60 normal; 4 SRA studies + TCGA)", f"{len(used)} = {n_t} tumour + {n_n} normal ({len(ss)} taken)", "n/a", "n/a", "n/a", "n/a",
             "Only 2 of the paper's 4 SRA studies are reproduced. T5_B (SRR13401190) excluded: R2 is a copy of R1 (also in the SRA run itself)."],
            ["Read-count agreement with S2 (runs)", len(rc), len(rc), n_rc_match, 0, len(rc) - n_rc_match, 0, "Our fastp raw read pairs vs paper Table S2 Total Reads; includes T5_B."]]
    def grow(lab, g, tail): return [lab, g["paper"], g["ours"], g["matched"], g["opposite"], g["missed"], g["extra"], tail + MS]
    for d in ("down", "up"):
        g = GS_fc[d]; rows.append(grow(f"{d.capitalize()} genes (featureCounts, primary)", g, f"Matched = significant in ours with the same direction. Opposite = paper-list genes significant in ours the other way ({g['opposite']}); {g['opposite_ours_side']} of our {d} genes are in the paper's other list. Symbol level, one row per symbol."))
    def allg(gs): return {k: gs["up"][k] + gs["down"][k] for k in gs["up"]}
    rows.append(grow("All DEGs (featureCounts, primary)", allg(GS_fc), "Sum of up and down."))
    for d in ("down", "up"):
        g = GS_sal[d]; rows.append(grow(f"{d.capitalize()} genes (Salmon, secondary)", g, f"Opposite: {g['opposite']} paper-list genes; {g['opposite_ours_side']} of our {d} genes are in the paper's other list."))
    rows.append(grow("All DEGs (Salmon, secondary)", allg(GS_sal), "Sum of up and down."))
    go = [PC[("GO", d)] for d in ("up", "down")]; S = lambda k: sum(x[k] for x in go)
    rows.append(["GO BP terms (up + down, padj < 0.05)", S("paper"), S("ours"), S("matched"), S("opposite"), S("missed"), S("extra"), "Matched by GO ID within the same direction (paper S4/S5 vs our clusterProfiler ORA). Biological process only."])
    for d in ("down", "up"):
        k = PC[("KEGG", d)]; rows.append([f"KEGG {d} pathways (padj < 0.05)", k["paper"], k["ours"], k["matched"], k["opposite"], k["missed"], k["extra"], "Matched by pathway name. KEGG release difference is a likely explanation for paper-only terms; not verified."])
    rows.append(["18-gene table (significant / expected direction)", 18, f"significant {g18_sig}, expected direction {g18_dir}/18", g18_match, g18_opp, 18 - g18_match - g18_opp, 0,
                 "Matched = significant (padj < 0.05, |log2FC| >= 1) in the paper's direction. Direction 17/18: APOA4 is untestable; CXCL8, KRT6A/B, KRT17 have padj NA (direction only)."])
    for i, r in enumerate(rows, 2):
        label, paper_n, ours_n, m, o, ms, ex, note = r
        for j, v in enumerate([label, paper_n, ours_n, m, o, ms, ex], 1): ws.cell(row=i, column=j, value=v).border = BOX
        ws.cell(row=i, column=11, value=note).border = BOX
        if m != "n/a":
            fcell(ws, i, 8, f'=IFERROR(D{i}/B{i},"n/a")', ratio(m, paper_n), PCT).border = BOX; fcell(ws, i, 9, f'=IFERROR(D{i}/C{i},"n/a")', ratio(m, ours_n), PCT).border = BOX
            fcell(ws, i, 10, f'=IF(D{i}+E{i}+F{i}=B{i},"OK","CHECK")', "OK" if m + o + ms == paper_n else "CHECK").border = BOX
        else:
            for j in (8, 9, 10): ws.cell(row=i, column=j, value="n/a").border = BOX
        ws.cell(row=i, column=11).alignment = Alignment(wrap_text=True, vertical="top"); ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        for j in (2, 3, 4, 5, 6, 7): ws.cell(row=i, column=j).number_format = "#,##0"
    last = len(rows) + 1; r0 = last + 2
    ws.cell(row=r0, column=1, value="Enrichment of our significant genes in the paper's DEG lists (featureCounts; universe = tested genes with a symbol)").font = Font(bold=True)
    header(ws, r0 + 1, ["Test", "Our genes", "Paper genes in universe", "Overlap", "Expected", "Fold enrichment", "p (hypergeometric)", "Universe N"])
    for i, (lab, d) in enumerate((("ours down vs paper down", "down"), ("ours up vs paper up", "up")), r0 + 2):
        e = ENR[d]
        for j, v in enumerate([lab, e["ours"], e["paper_in_universe"], e["overlap"], round(e["expected"], 1), round(e["fold"], 2), (e["p"] if e["p"] > 1e-300 else "< 1e-300"), N], 1): ws.cell(row=i, column=j, value=v).border = BOX
        ws.cell(row=i, column=7).number_format = SCI
    rb = r0 + 6; ws.cell(row=rb, column=1, value="How to read this").font = Font(bold=True, size=12)
    box = ["Agreement with the paper is not independent validation: our 20 samples are part of the paper's data (its Table S2 lists them).",
           "Recall = matched / paper; 'Our hits also in paper list' = matched / ours (not a precision: the paper used metaSeq on 120 samples).",
           "Each gene/term row adds up: matched + opposite direction + missed = paper (column J checks this).",
           "The paper used 120 samples (our 19 are about 16 %), so the low recall for UP genes is likely explained by the smaller sample size (less power); likely, not proven.",
           "Ours-only genes are not necessarily false positives (different method, 120 samples in the paper). Direction agreement is the meaningful comparison:",
           "S3 lists no fold change, so only direction can be compared. See Gene_overlap (every gene), Opposite_genes (the 10 opposite-direction genes) and Not_in_table_causes.",
           BAL, "Pathway agreement depends on the KEGG/GO release and gene-set background, which differ between the paper and this pipeline (not verified)."]
    if static: box.append("This is the values copy: no formulas (percentages are stored as numbers). The main file with live formulas is Hosseini2023_vs_Pipeline_comparison.xlsx.")
    for i, t in enumerate(box, rb + 1):
        for j in range(1, 12): ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(row=i, column=1, value=t)
    widths(ws, [44, 30, 36, 11, 12, 14, 14, 10, 14, 18, 80]); ws.freeze_panes = "B2"

    # 2 Gene_overlap -----------------------------------------------------------
    ws = sheet("Gene_overlap"); header(ws, 1, ["gene", "paper direction", "our direction (sign of log2FC)", "our log2FC (unshrunk)", "our padj", "our baseMean", "category", "Salmon direction", "Salmon log2FC", "Salmon padj", "Salmon significant"])
    union = list(dict.fromkeys(list(pdir) + list(fc.index[fc.sig]))); rows_go = []
    for g in union:
        r = fc.loc[g] if g in fc.index else None; s = sal.loc[g] if g in sal.index else None
        rows_go.append([g, pdir.get(g, ""), (r["dir"] if r is not None else ""), (None if r is None or pd.isna(r.log2FoldChange) else float(r.log2FoldChange)), (None if r is None or pd.isna(r.padj) else float(r.padj)),
                        (None if r is None else float(r.baseMean)), (CATS_fc[g] if g in pdir else CAT_X), (s["dir"] if s is not None else ""), (None if s is None or pd.isna(s.log2FoldChange) else float(s.log2FoldChange)),
                        (None if s is None or pd.isna(s.padj) else float(s.padj)), ("" if s is None else ("Yes" if s.sig else "No"))])
    order = [CAT_M, CAT_O] + CAT_P + [CAT_X]; rows_go.sort(key=lambda x: (order.index(x[6]), x[0]))
    for i, r in enumerate(rows_go, 2):
        for j, v in enumerate(r, 1): ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=4).number_format = L2; ws.cell(row=i, column=5).number_format = SCI; ws.cell(row=i, column=6).number_format = "#,##0.0"; ws.cell(row=i, column=9).number_format = L2; ws.cell(row=i, column=10).number_format = SCI
    ws.auto_filter.ref = f"A1:K{len(rows_go) + 1}"; ws.freeze_panes = "A2"
    df_go = pd.DataFrame(rows_go, columns=["gene", "pdir", "odir", "lfc", "padj", "bm", "cat", "sd", "sl", "sp", "ss"])
    header(ws, 1, ["category (featureCounts primary)", "all", "paper down", "paper up"], col0=13)
    for i, c in enumerate(order, 2):
        sub = df_go[df_go.cat == c]; ws.cell(row=i, column=13, value=c); ws.cell(row=i, column=14, value=len(sub)); ws.cell(row=i, column=15, value=int((sub.pdir == "down").sum())); ws.cell(row=i, column=16, value=int((sub.pdir == "up").sum()))
    i = len(order) + 2; ws.cell(row=i, column=13, value="Total genes in union").font = Font(bold=True); fcell(ws, i, 14, f"=SUM(N2:N{i - 1})", len(df_go))
    ws.cell(row=i + 2, column=13, value="One row per gene symbol; 'not in our DESeq2 table' includes symbol-alias mismatches (not hand-fixed); see Not_in_table_causes. Salmon columns are for the same symbols (secondary).")
    widths(ws, [14, 14, 18, 14, 12, 12, 60, 14, 13, 12, 12, 3, 62, 9, 12, 10])

    # 3 Missed_breakdown ---------------------------------------------------------
    ws = sheet("Missed_breakdown"); header(ws, 1, ["Category", "Down genes (n)", "Down % of paper list", "Up genes (n)", "Up % of paper list"])
    cats_pl = [CAT_M, CAT_O] + CAT_P; t = len(cats_pl) + 2
    cn = {c: (int(((df_go.cat == c) & (df_go.pdir == "down")).sum()), int(((df_go.cat == c) & (df_go.pdir == "up")).sum())) for c in cats_pl}; td, tu = sum(v[0] for v in cn.values()), sum(v[1] for v in cn.values())
    for i, c in enumerate(cats_pl, 2):
        ws.cell(row=i, column=1, value=c); ws.cell(row=i, column=2, value=cn[c][0]); ws.cell(row=i, column=4, value=cn[c][1])
        fcell(ws, i, 3, f"=B{i}/B${t}", cn[c][0] / td, PCT); fcell(ws, i, 5, f"=D{i}/D${t}", cn[c][1] / tu, PCT)
    ws.cell(row=t, column=1, value="Paper list size (total)").font = Font(bold=True)
    fcell(ws, t, 2, f"=SUM(B2:B{t - 1})", td); fcell(ws, t, 4, f"=SUM(D2:D{t - 1})", tu); fcell(ws, t, 3, f"=B{t}/B{t}", 1.0, PCT); fcell(ws, t, 5, f"=D{t}/D{t}", 1.0, PCT)
    ws.cell(row=t + 2, column=1, value="Reading: the share of the paper's list that we match, and for the rest the reason we miss it (padj, fold-change cutoff, Cook's flag, or gene absent from our table; causes for the last group are on Not_in_table_causes).")
    ws.cell(row=t + 3, column=1, value="Why the paper's up list is missed more often is likely the smaller sample size (19 vs 120 samples); this is a likely explanation, not proven.")
    ws.cell(row=t + 4, column=1, value=BAL)
    ch = BarChart(); ch.type = "col"; ch.title = "Paper DEGs: matched vs missed (share of paper list)"; ch.y_axis.title = "% of paper list"; ch.y_axis.numFmt = "0%"; ch.height = 9; ch.width = 22
    for col in (3, 5): ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=len(cats_pl) + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(cats_pl) + 1)); ws.add_chart(ch, f"A{t + 7}"); widths(ws, [64, 16, 20, 14, 20])

    # 4 Opposite_genes -----------------------------------------------------------
    ws = sheet("Opposite_genes"); header(ws, 1, ["gene", "paper direction", "our direction", "our log2FC (unshrunk)", "our padj", "our baseMean", "Salmon direction", "Salmon log2FC", "Salmon padj", "Salmon significant", "paper metaSeq", "paper p-value"])
    for i, g in enumerate(sorted(opp_genes, key=lambda x: fc.loc[x, "padj"]), 2):
        r = fc.loc[g]; s = sal.loc[g] if g in sal.index else None
        for j, v in enumerate([g, pdir[g], r["dir"], float(r.log2FoldChange), float(r.padj), float(r.baseMean), (s["dir"] if s is not None else ""), (None if s is None else float(s.log2FoldChange)), (None if s is None or pd.isna(s.padj) else float(s.padj)),
                               ("" if s is None else ("Yes" if s.sig else "No")), float(s3i.loc[g, "metaSeq"]), float(s3i.loc[g, "pvalue"])], 1): ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=4).number_format = L2; ws.cell(row=i, column=5).number_format = SCI; ws.cell(row=i, column=6).number_format = "#,##0.0"; ws.cell(row=i, column=8).number_format = L2; ws.cell(row=i, column=9).number_format = SCI; ws.cell(row=i, column=12).number_format = SCI
    ws.cell(row=len(opp_genes) + 3, column=1, value="Genes significant in our featureCounts analysis (padj < 0.05, |log2FC| >= 1) in the opposite direction to the paper's S3 list (paper: up in cancer, ours: down). Not explained here.")
    ws.freeze_panes = "A2"; widths(ws, [14, 15, 14, 18, 12, 12, 15, 14, 12, 16, 14, 14])

    # 5 Not_in_table_causes ---------------------------------------------------------
    ws = sheet("Not_in_table_causes"); header(ws, 1, ["gene", "paper direction", "cause", "our gene id(s)", "gene biotype", "max count in any sample", "samples with >= 10 counts"])
    for i, r in enumerate(notin, 2):
        for j, v in enumerate(r, 1): ws.cell(row=i, column=j, value=v)
    ws.auto_filter.ref = f"A1:G{len(notin) + 1}"; ws.freeze_panes = "A2"
    header(ws, 1, ["cause (paper genes missing from our DESeq2 table)", "paper down", "paper up", "all"], col0=9)
    for i, c in enumerate(CAUSES, 2):
        ws.cell(row=i, column=9, value=c); ws.cell(row=i, column=10, value=cause_cnt[c]["down"]); ws.cell(row=i, column=11, value=cause_cnt[c]["up"]); fcell(ws, i, 12, f"=J{i}+K{i}", cause_cnt[c]["down"] + cause_cnt[c]["up"])
    i = len(CAUSES) + 2; ws.cell(row=i, column=9, value="Total").font = Font(bold=True)
    for col, v in ((10, sum(v["down"] for v in cause_cnt.values())), (11, sum(v["up"] for v in cause_cnt.values())), (12, len(notin))):
        L = get_column_letter(col); fcell(ws, i, col, f"=SUM({L}2:{L}{i - 1})", v)
    ws.cell(row=i + 2, column=9, value="Low-count filter = our pre-filter (>= 10 counts in >= 5 of the 19 samples in the featureCounts matrix), applied before DESeq2 (as in nextflow/bin/deseq2.R).")
    ws.cell(row=i + 3, column=9, value="'No matching gene symbol': the paper's symbol is not a gene_name in our GENCODE v50 annotation (old/alias symbol); not hand-corrected.")
    ws.cell(row=i + 4, column=9, value="Biotype is the GENCODE gene_type of the gene id with the highest total count when a symbol has several ids.")
    widths(ws, [14, 15, 66, 26, 30, 20, 20, 3, 66, 12, 10, 8])

    # 6 Pathways -------------------------------------------------------------------
    ws = sheet("Pathways_GO_KEGG"); header(ws, 1, ["Type", "Direction", "Paper term", "GO / KEGG ID", "Paper adj. p", "In our results?", "Our term", "Our adjusted p", "Note"])
    rank = {"Yes": 0, "Opposite direction": 1, "No (paper only)": 2, "Ours only": 3}; pr = sorted(PATH_ROWS, key=lambda r: (r[0], r[1], rank[r[5]], r[4] if r[4] is not None else (r[7] or 1)))
    for i, r in enumerate(pr, 2):
        for j, v in enumerate(r, 1): ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=5).number_format = SCI; ws.cell(row=i, column=8).number_format = SCI
    ws.auto_filter.ref = f"A1:I{len(pr) + 1}"; ws.freeze_panes = "A2"; widths(ws, [8, 10, 60, 13, 13, 18, 60, 14, 90])

    # 7 Readcounts / 8 Gene18 / 9 Method_notes --------------------------------------
    ws = sheet("Readcounts_S2"); header(ws, 1, ["Run", "Paper total reads (S2)", "Our total reads (fastp pairs)", "Match", "Label", "Study", "Group", "Note"])
    for i, r in enumerate(rc, 2):
        for j, v in enumerate(r, 1): ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=2).number_format = "#,##0"; ws.cell(row=i, column=3).number_format = "#,##0"
    ws.auto_filter.ref = f"A1:H{len(rc) + 1}"; ws.freeze_panes = "A2"; widths(ws, [14, 22, 26, 9, 8, 12, 9, 48])
    ws = sheet("Gene18"); header(ws, 1, ["gene", "paper direction (S3)", "our direction", "our padj", "our log2FC (unshrunk)", "significant?", "significant in paper's direction?", "caveat", "Salmon direction", "Salmon padj"])
    for i, r in enumerate(g18, 2):
        for j, v in enumerate(r, 1): ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=4).number_format = SCI; ws.cell(row=i, column=5).number_format = L2; ws.cell(row=i, column=10).number_format = SCI
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:J{len(g18) + 1}"; widths(ws, [12, 18, 14, 12, 18, 13, 22, 90, 15, 12])
    ws = sheet("Method_notes"); header(ws, 1, ["Item", "Value"])
    tv = pd.read_csv(INPUTS["tool versions"], sep="\t").set_index("tool"); summ = json.load(open(INPUTS["ours DESeq2 summary"]))
    ver = lambda k: str(tv.loc[k, "version_reported_by_tool"]).strip()
    mrows = [("File version", "VALUES COPY (no formulas)" if static else "Live formulas (percentages and totals are Excel formulas; they calculate when opened in Excel). A values-only copy is Hosseini2023_vs_Pipeline_comparison_values.xlsx"),
             ("Aligner / quantifiers / QC", f"fastp {ver('fastp')}; STAR {ver('STAR')}; featureCounts {ver('featureCounts')}; Salmon {ver('salmon')}; samtools {ver('samtools')}; MultiQC {ver('multiqc')}"),
             ("DE model", f"DESeq2 {ver('R:DESeq2')} ({ver('R')}), design {summ['design']}, reference {summ['reference']}, apeglm shrinkage for plots only; tximport {ver('R:tximport')} for the Salmon route"),
             ("Pathways", f"clusterProfiler {ver('R:clusterProfiler')} ORA (GO BP; KEGG REST), set size 10-500, BH q 0.05; fgsea {ver('R:fgsea')}"),
             ("Reference", "hg38 primary chromosomes (chr1-22, X, Y, M); " + str(tv.loc["GENCODE annotation", "version_reported_by_tool"])),
             ("Significance rule", f"padj < {summ['padj_cut']} and |log2FC| >= {summ['lfc_cut']} on the UNSHRUNK log2FC (as in nextflow/bin/deseq2.R); one row per gene symbol (highest baseMean among tested rows)"),
             ("Low-count pre-filter", ">= 10 counts in >= 5 of the 19 samples (featureCounts matrix); genes failing it are not in the DESeq2 table"),
             ("Samples", f"{len(used)} used ({n_t} tumour, {n_n} normal); excluded: " + ", ".join(sorted(set(ss.label) - set(used.label))) + " (T5_B: R2 is a copy of R1)"),
             ("Sample list", "; ".join(f"{r.label}={r.srr}" for r in used.itertuples())),
             ("Paper", "Hosseini & Nemati, Sci Rep 2023, doi 10.1038/s41598-023-44459-y; supplement docs/reference/hosseini_2023_MOESM2_ESM.xlsx (md5 d260d2c65ef0293586f8b60792509de9)"),
             ("Paper vs ours", "Paper: 4 SRA studies + TCGA READ, 120 samples, metaSeq. Ours: 2 SRA studies (SRP219837, SRP301216), 19 samples, DESeq2. S3 has no fold change: direction only."),
             ("Sheets Summary / Gene_overlap / Missed_breakdown / Opposite_genes", f"{INPUTS['paper S3 (DEGs, up/down)']}; {INPUTS['ours DESeq2 featureCounts (primary)']}; {INPUTS['ours DESeq2 Salmon (secondary)']}"),
             ("Sheet Not_in_table_causes", f"{INPUTS['ours gene annotation (id, symbol, biotype)']}; {INPUTS['ours featureCounts count matrix (19 samples)']}; {INPUTS['ours DESeq2 summary']}"),
             ("Sheet Pathways_GO_KEGG", "; ".join(INPUTS[k] for k in INPUTS if k.startswith(("paper S4", "paper S5", "paper S6", "ours ORA")))),
             ("Sheet Readcounts_S2", f"{INPUTS['paper S2 (read counts)']}; {INPUTS['ours fastp JSON (per run)']}; {INPUTS['ours sample map']}"),
             ("Sheet Gene18", f"{INPUTS['paper S3 (DEGs, up/down)']}; {INPUTS['ours DESeq2 featureCounts (primary)']}; {INPUTS['step4b leave-one-out T2_B']}; {INPUTS['step8 refit without N4_A,N5_A']}"),
             ("Build script", f"{OUT}/build_comparison.py (pandas {pd.__version__}, openpyxl)")]
    for i, (a, b) in enumerate(mrows, 2):
        ws.cell(row=i, column=1, value=a).alignment = Alignment(vertical="top"); ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, [44, 150]); ws.freeze_panes = "A2"
    os.makedirs(OUT, exist_ok=True); wb.save(path); print(f"written {path} ({os.path.getsize(path)} bytes)")

build(False, XLSX); build(True, XLSX_VAL)
